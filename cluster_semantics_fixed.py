#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Скрипт кластеризации семантики для рекламного агентства
Гео: Калуга
Автор: SEO-специалист
"""

import re
from collections import defaultdict

# Конфигурация
INPUT_FILE = '/workspace/рекламное агенство.csv'
OUTPUT_FILE = '/workspace/semantics_kaluga_clustered.xlsx'

def is_garbage(phrase):
    """Проверяет, является ли запрос мусорным"""
    phrase_lower = phrase.lower().strip()
    
    # Целевые запросы - никогда не мусор (если нет других городов)
    if 'рекламное агенство' in phrase_lower or 'рекламное агентство' in phrase_lower:
        # Проверяем на другие города и мусор
        garbage_markers = [
            'москв', 'спб', 'петербург', 'казан', 'екатеринбург', 'новосибирск',
            'челябинск', 'самар', 'омск', 'ростов', 'уф', 'краснодар', 'воронеж',
            'перм', 'волгоград', 'саратов', 'тюмень', 'ижевск', 'барнаул',
            'иркутск', 'хабаровск', 'ярославль', 'владивосток', 'томск', 'оренбург',
            'кемерово', 'брянск', 'тверь', 'пенза', 'липецк', 'тула', 'киров',
            'курск', 'минск', 'алматы', 'астана', 'украина', 'киев', 'крым',
            'луганск', 'донецк', 'чита', 'абакан', 'белгород', 'вологда',
            'иваново', 'кострома', 'смоленск', 'тамбов', 'орел', 'псков',
            'сериал', 'фильм', 'кино', 'актер', 
            'что', 'это', 'как', 'где', 'когда', 'почему', 'зачем',
            'ооо', 'ип', 'ао', 'пао', 'оквэд', 'егрюл', 'инн', 'кпп', 'огрн',
        ]
        for marker in garbage_markers:
            if marker in phrase_lower and marker not in {'калуг', 'обнинск'}:
                return True
        return False
    
    return True  # Все что не содержит "рекламное агенство/агентство" - мусор


def get_clusters(phrases_data):
    """Группирует фразы по кластерам"""
    clusters = defaultdict(list)
    
    for item in phrases_data:
        phrase = item['phrase']
        phrase_lower = phrase.lower()
        
        # Определяем кластер
        if 'калуг' in phrase_lower or 'обнинск' in phrase_lower:
            if 'наружн' in phrase_lower or 'вывеск' in phrase_lower or 'баннер' in phrase_lower or 'щит' in phrase_lower:
                clusters['Наружная реклама Калуга'].append(item)
            elif 'полиграфи' in phrase_lower or 'листовк' in phrase_lower or 'визитк' in phrase_lower or 'печат' in phrase_lower:
                clusters['Полиграфия Калуга'].append(item)
            elif 'широкоформатн' in phrase_lower or 'пленк' in phrase_lower:
                clusters['Широкоформатная печать Калуга'].append(item)
            elif 'pos' in phrase_lower or 'wobbler' in phrase_lower or 'стоппер' in phrase_lower:
                clusters['POS-материалы Калуга'].append(item)
            elif 'сувенир' in phrase_lower or 'мерч' in phrase_lower or 'promo' in phrase_lower:
                clusters['Сувенирная продукция Калуга'].append(item)
            elif 'дизайн' in phrase_lower or 'логотип' in phrase_lower or 'брендбук' in phrase_lower:
                clusters['Дизайн и бренд Калуга'].append(item)
            else:
                clusters['Общие запросы Калуга'].append(item)
        else:
            # Без гео - общие запросы
            clusters['Общие запросы (Россия)'].append(item)
    
    return dict(clusters)


def main():
    print("🚀 Начинаем обработку семантики для рекламного агентства (Калуга)")
    print("=" * 70)
    
    # Чтение данных из CSV
    phrases_data = []
    with open(INPUT_FILE, 'r', encoding='utf-8-sig') as f:
        content = f.read().replace('\r\n', '\n').replace('\r', '\n')
        lines = content.split('\n')
        
        header_line = lines[0]
        header = header_line.split(';')
        print(f"Заголовки: {header}")
        
        for line in lines[1:]:
            if not line.strip():
                continue
            row = line.split(';')
            if len(row) >= 5:
                phrase = row[0].strip('"').strip()
                try:
                    frequency_all = int(row[3]) if row[3] else 0
                    frequency_exact = int(row[4]) if row[4] else 0
                except ValueError:
                    frequency_all = 0
                    frequency_exact = 0
                
                phrases_data.append({
                    'phrase': phrase,
                    'frequency_all': frequency_all,
                    'frequency_exact': frequency_exact
                })
    
    print(f"✅ Загружено фраз: {len(phrases_data)}")
    
    # Фильтрация мусора
    clean_phrases = []
    garbage_phrases = []
    
    for item in phrases_data:
        phrase = item['phrase']
        if is_garbage(phrase):
            garbage_phrases.append(item)
        else:
            clean_phrases.append(item)
    
    print(f"🗑️ Отфильтровано мусора: {len(garbage_phrases)}")
    print(f"✨ Осталось чистых фраз: {len(clean_phrases)}")
    
    # Приоритет фразам с Калугой
    kaluga_phrases = [p for p in clean_phrases if 'калуг' in p['phrase'].lower()]
    other_phrases = [p for p in clean_phrases if 'калуг' not in p['phrase'].lower()]
    
    print(f"📍 Фраз с Калугой: {len(kaluga_phrases)}")
    print(f"📄 Остальных фраз: {len(other_phrases)}")
    
    # Кластеризация
    clusters = get_clusters(kaluga_phrases + other_phrases)
    
    print(f"\n📊 Получено кластеров: {len(clusters)}")
    for cluster_name, items in clusters.items():
        print(f"  • {cluster_name}: {len(items)} фраз")
    
    # Сохранение в Excel
    try:
        from openpyxl import Workbook
        import openpyxl
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Семантика_Калуга"
        
        headers = ['Кластер', 'Подкластер', 'Ключевая фраза', 'Частотность (все)', 'Частотность (!!)', 'Гео', 'Статус очистки']
        ws.append(headers)
        
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = openpyxl.styles.Font(color="FFFFFF", bold=True)
        
        for cluster_name, items in clusters.items():
            for item in items:
                geo = 'Калуга' if 'калуг' in item['phrase'].lower() or 'обнинск' in item['phrase'].lower() else 'Россия'
                ws.append([
                    cluster_name, 
                    '-', 
                    item['phrase'], 
                    item['frequency_all'], 
                    item['frequency_exact'], 
                    geo, 
                    'Сохранен'
                ])
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(OUTPUT_FILE)
        print(f"\n💾 Файл сохранен: {OUTPUT_FILE}")
        
    except ImportError:
        print("\n⚠️ Модуль openpyxl не установлен.")
        output_csv = OUTPUT_FILE.replace('.xlsx', '.csv')
        with open(output_csv, 'w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f, delimiter=';')
            writer.writerow(['Кластер', 'Подкластер', 'Ключевая фраза', 'Частотность (все)', 'Частотность (!!)', 'Гео', 'Статус очистки'])
            
            for cluster_name, items in clusters.items():
                for item in items:
                    geo = 'Калуга' if 'калуг' in item['phrase'].lower() else 'Россия'
                    writer.writerow([cluster_name, '-', item['phrase'], item['frequency_all'], item['frequency_exact'], geo, 'Сохранен'])
        
        print(f"💾 Файл сохранен (CSV): {output_csv}")
    
    # Вывод примеров
    print("\n" + "=" * 70)
    print("📋 Примеры чистых фраз для Калуги:")
    print("=" * 70)
    for i, item in enumerate(kaluga_phrases[:15], 1):
        print(f"{i:2}. {item['phrase']} (частотность: {item['frequency_all']})")
    
    print("\n" + "=" * 70)
    print("🎯 Кластеризация завершена успешно!")
    print("=" * 70)


if __name__ == '__main__':
    main()
